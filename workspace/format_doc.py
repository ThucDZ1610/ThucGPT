import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc = docx.Document()

# Adjust default style
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(12)

with open(r'c:\cac\chitiet.docx', 'r', encoding='utf-8') as f:
    lines = f.readlines()

for line in lines:
    line = line.strip('\n')
    if not line:
        doc.add_paragraph()
        continue
    
    # Check for headings
    heading_level = 0
    if line.startswith('# '): heading_level = 1
    elif line.startswith('## '): heading_level = 2
    elif line.startswith('### '): heading_level = 3
    
    if heading_level > 0:
        p = doc.add_heading(level=heading_level)
        # Keep the exact text as requested, just styled
        run = p.add_run(line)
        run.font.name = 'Calibri'
        if heading_level == 1:
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 153) # Dark blue
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif heading_level == 2:
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 102, 204)
        elif heading_level == 3:
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 102, 204)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        # Handle bold formatting like **text**
        # We will keep the asterisks to strictly obey "do not delete" but make them bold.
        # Actually, if I keep asterisks and make it bold, it looks weird. 
        # I'll just add the whole line, but check if there are ** to make parts bold.
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part)
                run.bold = True
            else:
                p.add_run(part)
        
        # Adjust spacing
        p.paragraph_format.space_after = Pt(6)

doc.save(r'c:\cac\chitiet_formatted.docx')
